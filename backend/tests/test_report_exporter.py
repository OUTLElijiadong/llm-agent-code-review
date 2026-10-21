"""T11 报告导出服务单元测试

覆盖范围:
- JSON / dict / HTML 三种导出形态的结构与内容
- 3 套 Jinja2 模板(simple/detailed/compliance)渲染与关键字段
- 按严重度分组(严重>高>中>低)、按类型分组(9 种 issue_type)
- 合规摘要聚合(4 套标准)、Top N 漏洞按 CVSS 降序
- 空问题列表、单问题、多问题混合场景
- IssueOut Pydantic 模型与原生 dict 输入兼容性
- 内置模板载入
"""
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html.parser import HTMLParser
from typing import Any, Dict, List

import pytest
from docx import Document
from docx.oxml.ns import qn
from jinja2.exceptions import SecurityError

from app.schemas.review import IssueOut
from app.services import report_pdf_exporter
from app.services.report_exporter import (
    _build_compliance_summary,
    _build_report_context,
    _build_top_vulnerabilities,
    _group_issues_by_severity,
    _group_issues_by_type,
    export_to_dict,
    export_to_html,
    export_to_json,
    finding_status_label,
    human_review_status_label,
    load_builtin_template,
    report_status_label,
)
from app.services.report_pdf_exporter import (
    _build_header_elements,
    _build_issue_section,
    _build_styles,
    _ensure_chinese_font,
    export_to_pdf,
)
from app.services.report_word_exporter import _format_score as format_word_score
from app.services.report_word_exporter import export_to_word


@pytest.mark.parametrize(
    ("labeler", "raw_value", "expected"),
    [
        (report_status_label, "unexpected-task-state", "状态未识别"),
        (finding_status_label, "unexpected-finding-state", "状态未识别"),
        (finding_status_label, "refuted", "已证伪"),
        (human_review_status_label, "unexpected-review-state", "状态未识别"),
        (report_status_label, None, "未记录"),
        (finding_status_label, "", "未记录"),
    ],
)
def test_report_status_labels_have_safe_fallbacks(labeler, raw_value, expected):
    """未知与缺失枚举必须明确标注，不能伪装成已完成或泄漏原始值。"""
    assert labeler(raw_value) == expected

# ============ 测试数据工厂 ============

def _make_task(**overrides: Any) -> Dict[str, Any]:
    """构造测试用 task 字典。

    Args:
        **overrides: 覆盖默认字段的关键字参数。

    Returns:
        Dict[str, Any]: 含 task 基础字段的字典。
    """
    base: Dict[str, Any] = {
        "id": 1,
        "task_name": "示例审查任务",
        "project_id": 10,
        "project_name": "示例项目",
        "review_type": "security",
        "status": "success",
        "total_files": 3,
        "processed_files": 3,
        "total_issues": 5,
        "severe_issues": 1,
        "high_issues": 2,
        "medium_issues": 1,
        "low_issues": 1,
        "score": 72,
        "coverage": {"coverage_status": "complete", "score_available": True},
        "summary": "发现多处安全漏洞,建议尽快修复。",
        "model_name": "deepseek-v3",
        "duration_ms": 5600,
        "start_time": "2026-06-25T10:00:00",
        "end_time": "2026-06-25T10:05:36",
        "create_time": "2026-06-25T10:05:36",
    }
    base.update(overrides)
    return base


def _make_issue(**overrides: Any) -> Dict[str, Any]:
    """构造测试用 issue 字典。

    Args:
        **overrides: 覆盖默认字段的关键字参数。

    Returns:
        Dict[str, Any]: 含 issue 全量字段的字典。
    """
    base: Dict[str, Any] = {
        "id": 1,
        "task_id": 1,
        "file_id": 100,
        "file_name": "app.py",
        "line_number": 42,
        "end_line": 45,
        "issue_type": "安全漏洞",
        "severity": "高",
        "title": "SQL 注入漏洞",
        "description": "用户输入未经过滤直接拼接进 SQL 查询,可能导致数据泄露。",
        "suggestion": "使用参数化查询替代字符串拼接。",
        "fixed_code": "cursor.execute(\"SELECT * FROM users WHERE id=%s\", (user_id,))",
        "status": "unfixed",
        "create_time": "2026-06-25T10:05:00",
        "owasp": "A03:2021-Injection",
        "cwe": "CWE-89",
        "evidence": "query = \"SELECT * FROM users WHERE id=\" + user_id",
        "exploit_scenario": "攻击者构造 user_id=1 OR 1=1 可绕过认证获取全部用户数据。",
        "references_json": ["https://owasp.org/www-community/attacks/SQL_Injection"],
        "confidence": 0.95,
        "source": "llm",
        "source_details": [
            {"source": "llm:reviewer", "confidence": 0.95},
            {"source": "static:security", "confidence": 1.0},
        ],
        "confirmation_count": 2,
        "finding_fingerprint": "finding-fingerprint-001",
        "cvss_score": 8.6,
        "cvss_vector": "AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
        "cvss_version": "3.1",
        "cvss_source": "vector",
        "compliance_mapping": {
            "iso27001": ["A.8.25", "A.8.28"],
            "gdpr": ["Art.32"],
            "pci_dss": ["Req-6.4.1"],
            "hipaa": ["§164.312(a)"],
        },
        "remediation": "1. 改用参数化查询;2. 增加输入校验;3. 最小化数据库账号权限。",
        "static_rule_hits": 2,
        "handled_by": None,
        "handled_at": None,
        "update_time": "2026-06-25T10:06:00",
    }
    base.update(overrides)
    return base


def _make_mixed_issues() -> List[Dict[str, Any]]:
    """构造覆盖 4 种严重度 + 9 种 issue_type 的混合 issue 列表。

    Returns:
        List[Dict[str, Any]]: 9 条 issue,覆盖严重/高/中/低 与 9 种类型。
    """
    types = ["安全漏洞", "潜在Bug", "性能问题", "代码风格", "可维护性",
             "兼容性", "并发问题", "资源泄漏", "配置问题"]
    severities = ["严重", "高", "高", "中", "中", "低", "低", "严重", "高"]
    cwes = ["CWE-89", "CWE-79", "CWE-78", "CWE-22", "CWE-200",
            "CWE-327", "CWE-502", "CWE-798", "CWE-918"]
    cvss = [9.1, 8.6, 8.2, 6.5, 5.4, 3.2, 7.1, 9.8, 8.0]
    issues: List[Dict[str, Any]] = []
    for idx in range(9):
        issues.append(_make_issue(
            id=idx + 1,
            issue_type=types[idx],
            severity=severities[idx],
            title=f"{types[idx]} #{idx + 1}",
            description=f"{types[idx]} 详细描述。",
            cwe=cwes[idx],
            cvss_score=cvss[idx],
            compliance_mapping=None,  # 触发按 cwe 自动补全
            file_name=f"file_{idx + 1}.py",
            line_number=idx * 10 + 1,
        ))
    return issues


class _LinkCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.hrefs: List[str] = []
        self.active_tags: List[str] = []

    def handle_starttag(self, tag: str, attrs: List[tuple[str, str | None]]) -> None:
        if tag in {"script", "svg", "iframe"}:
            self.active_tags.append(tag)
        for name, value in attrs:
            if name == "href" and value is not None:
                self.hrefs.append(value)
            if name.startswith("on"):
                self.active_tags.append(name)


@pytest.mark.parametrize("template_type", ["detailed", "compliance", "historical"])
def test_html_report_references_reject_active_schemes_including_historical_templates(template_type: str) -> None:
    """报告中的模型引用只能生成可点击的 HTTP(S) 链接。"""
    issue = _make_issue(
        title='<svg onload="alert(1)">',
        references_json=[
            "https://example.com/safe",
            "http://example.org/reference",
            "javascript:alert(1)",
            "JaVaScRiPt:alert(1)",
            "data:text/html,<script>alert(1)</script>",
            "java\nscript:alert(1)",
            "java\tscript:alert(1)",
            " javascript:alert(1)",
            "//evil.example/path",
            "https://",
            "<svg onload=alert(1)>",
        ],
    )
    template = (
        '{% for ref in issues[0].references_json %}<a href="{{ ref }}">{{ ref }}</a>{% endfor %}'
        if template_type == "historical" else load_builtin_template(template_type)
    )
    html = export_to_html(_make_task(), [issue], "摘要", 80, template)
    links = _LinkCollector()
    links.feed(html)

    assert links.hrefs == (["#issue-1"] if template_type == "detailed" else []) + [
        "https://example.com/safe", "http://example.org/reference",
    ]
    assert links.active_tags == []
    if template_type != "historical":
        assert "&lt;svg" in html  # 审查标题保留为文本，不生成活动节点


def test_html_reference_filter_does_not_mutate_json_audit_evidence() -> None:
    issue = _make_issue(references_json=["javascript:alert(1)", "https://example.com/safe"])
    export_to_html(_make_task(), [issue], "摘要", 80, load_builtin_template("detailed"))

    assert issue["references_json"] == ["javascript:alert(1)", "https://example.com/safe"]
    assert export_to_dict(_make_task(), [issue], "摘要", 80)["issues"][0]["references_json"] == issue["references_json"]


def test_report_context_removes_ai_markers_but_preserves_code() -> None:
    issue = _make_issue(
        title="**输入校验**问题",
        description="*需要检查* 用户输入。",
        suggestion="**先校验，再处理。**",
        fixed_code="if value ** 2 > limit:\n    reject()",
    )

    payload = json.loads(export_to_json(_make_task(), [issue], "**总体结论**，*需要复核*。", 80))

    assert payload["summary"] == "总体结论，需要复核。"
    assert payload["issues"][0]["title"] == "输入校验问题"
    assert payload["issues"][0]["description"] == "需要检查 用户输入。"
    assert payload["issues"][0]["suggestion"] == "先校验，再处理。"
    assert payload["issues"][0]["fixed_code"] == "if value ** 2 > limit:\n    reject()"


def test_html_report_top_vulnerability_alias_and_ai_text_are_safe() -> None:
    issue = _make_issue(
        description='<img src=x onerror="alert(1)">',
        references_json=["javascript:alert(1)", "https://example.com/safe"],
    )
    template = (
        '<div>{{ summary }}</div><div>{{ issues[0].description }}</div>'
        '{% for ref in statistics.top_vulnerabilities[0].references_json %}'
        '<a href="{{ ref }}">{{ ref }}</a>{% endfor %}'
    )
    html = export_to_html(_make_task(), [issue], '<svg onload="alert(1)">', 80, template)
    links = _LinkCollector()
    links.feed(html)

    assert links.hrefs == ["https://example.com/safe"]
    assert links.active_tags == []
    assert "&lt;svg" in html and "&lt;img" in html


def test_concurrent_html_exports_keep_sanitizer_state_isolated() -> None:
    template = '<div>{{ task_info.task_name }}</div><script>window.bad=1</script>'

    def render(index: int) -> str:
        return export_to_html(_make_task(task_name=f"report-{index}"), [], "", 100, template)

    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(render, range(32)))

    for index, html in enumerate(results):
        assert f"report-{index}" in html
        assert "<script" not in html
        assert "window.bad" in html  # 活动标签被剥离，证据文本仍可读


def test_custom_template_safe_filter_cannot_reactivate_ai_html() -> None:
    html = export_to_html(
        _make_task(), [], '<svg onload="alert(1)"></svg>', 100,
        '<div>{{ summary | safe }}</div><img src="data:text/html,attack" onerror="alert(1)">',
    )
    links = _LinkCollector()
    links.feed(html)

    assert links.active_tags == []
    assert "<img" in html and "onerror=" not in html and "data:text/html" not in html


def test_html_report_does_not_autoload_external_images_or_fonts() -> None:
    """旧自定义模板的远程图片与 CSS 资源不应随报告打开而自动加载。"""
    template = (
        '<style>.report-proof{background-image:url(https://c08-probe.example/css)}'
        '@font-face{font-family:probe;src:url(https://c08-probe.example/font)}</style>'
        '<div class="report-proof" style="color:rgb(12,34,56)">静态样式</div>'
        '<img src="https://c08-probe.example/image" alt="remote">'
        '<img src="data:image/png;base64,iVBORw0KGgo=" alt="inline">'
    )
    html = export_to_html(_make_task(), [], "", 100, template)

    assert "img-src data:; font-src data:" in html
    assert 'src="https://c08-probe.example/image"' not in html
    assert 'src="data:image/png;base64,iVBORw0KGgo="' in html
    assert ".report-proof{background-image:" in html  # CSS 样式可编辑，由 CSP 阻止外部加载。
    assert 'style="color:rgb(12,34,56)"' in html


# ============ JSON / dict 导出结构 ============

def test_export_to_json_returns_parseable_json_with_required_keys():
    """JSON 导出应可被 json.loads 解析,且包含 5 个顶层键。"""
    task = _make_task()
    issues = [_make_issue()]
    raw = export_to_json(task, issues, "摘要", 80)
    import json
    payload = json.loads(raw)
    assert set(["task_info", "summary", "score", "issues", "statistics"]).issubset(payload.keys())


def test_export_to_json_contains_task_info_and_score():
    """JSON 字符串应包含 task_name 与 score 文本。"""
    task = _make_task(task_name="JSON任务")
    raw = export_to_json(task, [], "摘要", 88)
    assert "JSON任务" in raw
    assert "88" in raw


def test_terminal_report_state_and_coverage_survive_all_standard_exports():
    task = _make_task(
        status="cancelled",
        error_message="用户停止审查",
        processed_files=1,
        total_files=3,
        coverage={
            "coverage_status": "partial",
            "completed_files": 1,
            "unscanned_files": 2,
            "excluded_files": ["vendor/generated.py"],
        },
    )
    json_payload = __import__("json").loads(export_to_json(task, [], "任务已取消", 0))
    html_exports = [
        export_to_html(task, [], "任务已取消", 0, load_builtin_template(template_type))
        for template_type in ("simple", "detailed", "compliance")
    ]
    pdf = export_to_pdf(task, [], "任务已取消", 0)
    word = export_to_word(task, [], "任务已取消", 0)

    assert json_payload["task_info"]["status"] == "cancelled"
    assert json_payload["task_info"]["error_message"] == "用户停止审查"
    assert json_payload["task_info"]["coverage"]["unscanned_files"] == 2
    assert all("已取消" in html for html in html_exports)
    assert all("cancelled" not in html for html in html_exports)
    assert all("用户停止审查" in html for html in html_exports)
    assert all("vendor/generated.py" in html for html in html_exports)
    assert pdf.startswith(b"%PDF")
    pdf_header = "\n".join(
        getattr(element, "text", "")
        for element in _build_header_elements(
            json_payload["task_info"], 0, _build_styles(_ensure_chinese_font())
        )
    )
    assert "已取消" in pdf_header
    assert "cancelled" not in pdf_header
    assert "用户停止审查" in pdf_header
    assert "扫描范围与覆盖情况" not in pdf_header
    pdf_text = "\n".join(export_to_dict(task, [], "任务已取消", 0)["scope_lines"])
    assert "vendor/generated.py" in pdf_text
    word_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs)
    assert "已取消" in word_text
    assert "cancelled" not in word_text
    assert "任务状态：已取消" in pdf_text
    assert "用户停止审查" in word_text
    assert "vendor/generated.py" in word_text


def test_issue_and_decompilation_states_are_localized_in_all_visual_exports(monkeypatch):
    """HTML/PDF/Word 对外展示的审查、复核、证据、冲突及反编译状态都使用中文。"""
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table

    issue = _make_issue(
        status="fixed",
        source="static",
        source_details=None,
        cvss_source="vector",
        aggregation_version="finding-aggregation-v1",
        evidence_quality="verified",
        conflict_status="unresolved",
        human_review_status="evidence_requested",
    )
    evidence = {"decompilation": {"status": "succeeded", "tool": "jadx"}}
    expected_labels = (
        "已修复", "静态规则", "向量计算", "证据引用已核对", "待处理冲突", "待补充证据", "执行成功",
    )
    raw_values = ("fixed", "static", "vector", "verified", "unresolved", "evidence_requested", "succeeded")
    html_outputs = [
        export_to_html(_make_task(), [issue], "", 72, load_builtin_template(template_type), evidence=evidence)
        for template_type in ("simple", "detailed", "compliance")
    ]
    assert all("执行成功" in html for html in html_outputs)
    assert all("已修复" in html for html in html_outputs)
    assert all(all(label in html for label in expected_labels[:-1]) for html in html_outputs[1:])
    assert "succeeded" not in html_outputs[0]
    assert all(not any(raw in html for raw in raw_values) for html in html_outputs[1:])

    original_build = SimpleDocTemplate.build
    pdf_visible_text = []

    def collect_flowable_text(flowable):
        if isinstance(flowable, Paragraph):
            pdf_visible_text.append(flowable.getPlainText())
        elif isinstance(flowable, Table):
            for row in flowable._cellvalues:
                for cell in row:
                    if isinstance(cell, (list, tuple)):
                        for nested in cell:
                            collect_flowable_text(nested)
                    else:
                        collect_flowable_text(cell)

    def capture_build(doc, flowables, *args, **kwargs):
        for flowable in flowables:
            collect_flowable_text(flowable)
        return original_build(doc, flowables, *args, **kwargs)

    monkeypatch.setattr(SimpleDocTemplate, "build", capture_build)
    pdf = export_to_pdf(_make_task(), [issue], "", 72, evidence=evidence)
    assert pdf.startswith(b"%PDF")
    pdf_text = "\n".join(pdf_visible_text)
    assert all(label in pdf_text for label in expected_labels)
    assert not any(raw in pdf_text for raw in raw_values)

    word = export_to_word(_make_task(), [issue], "", 72, evidence=evidence)
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        word_xml = archive.read("word/document.xml").decode("utf-8")
    assert all(label in word_xml for label in expected_labels)
    assert not any(raw in word_xml for raw in raw_values)


def test_scope_coverage_is_human_readable_in_html_word_and_pdf_exports(monkeypatch):
    """覆盖与来源对象应逐项展示，不能把内部 JSON 原样塞进读者报告。"""
    task = _make_task(
        processed_files=2,
        total_files=5,
        coverage={
            "coverage_status": "partial",
            "verification_status": "failed",
            "completed_files": 2,
            "unscanned_file_names": ["src/later.py", "src/queued.py"],
            "excluded_files": [
                {"file_name": "vendor/jquery.js", "reason": "第三方依赖"},
            ],
            "errors": ["扫描请求超时"],
            "scan_metadata": {"duration_seconds": 13},
        },
        source={
            "type": "remote_site",
            "stats_basis": "actual_task_events",
            "request_id": "req-coverage-123",
        },
    )

    issue = _make_issue(source_details=[
        {"source": "llm:reviewer", "confidence": 0.95},
        {"source": "static:security", "confidence": 1.0},
    ])
    htmls = [
        export_to_html(task, [issue], "", 72, load_builtin_template(template_type))
        for template_type in ("simple", "detailed", "compliance")
    ]
    word = export_to_word(task, [issue], "", 72)
    pdf_paragraphs: list[str] = []
    original_paragraph = report_pdf_exporter.Paragraph

    def record_pdf_paragraph(text, style, *args, **kwargs):
        pdf_paragraphs.append(str(text))
        return original_paragraph(text, style, *args, **kwargs)

    monkeypatch.setattr(report_pdf_exporter, "Paragraph", record_pdf_paragraph)
    assert export_to_pdf(task, [issue], "", 72).startswith(b"%PDF")
    word_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs)
    pdf_text = "\n".join(pdf_paragraphs)

    for rendered in (*htmls, word_text, pdf_text):
        assert "部分覆盖" in rendered
        assert "未通过" in rendered
        assert "覆盖状态：部分覆盖" in rendered
        assert "验证状态：未通过" in rendered
        assert "已完成文件数：2" in rendered
        assert "未扫描文件名称：src/later.py、src/queued.py" in rendered
        assert "排除文件第1项文件名：vendor/jquery.js" in rendered
        assert "排除文件第1项原因：第三方依赖" in rendered
        assert "错误：扫描请求超时" in rendered
        assert "扫描元数据持续秒数：13" in rendered
        assert "来源类型：remote_site" in rendered
        assert "统计依据：actual_task_events" in rendered
        assert "请求编号：req-coverage-123" in rendered
        assert '"coverage_status"' not in rendered
        assert '"stats_basis"' not in rendered
        assert '"source_details"' not in rendered

    for rendered in (htmls[1], htmls[2], word_text, pdf_text):
        assert "来源明细第1项来源：llm:reviewer" in rendered
        assert "来源明细第1项置信度：0.95" in rendered
        assert '"source":' not in rendered


def test_scope_exports_hide_internal_metadata_and_bound_large_payloads(monkeypatch):
    """大扫描元数据只保留已映射事实，不把内部字段扩成报告正文。"""
    task = _make_task(
        coverage={
            "coverage_status": "partial",
            "scan_metadata": {
                "duration_seconds": 13,
                "provider_payload": {"secret_marker": "INTERNAL_SECRET_1", "files": ["private.py"]},
                **{f"internal_field_{index}": f"INTERNAL_SECRET_{index}" for index in range(1_896)},
            },
        },
    )

    context = export_to_dict(task, [], "", 72)
    scope_text = "\n".join(context["scope_lines"])
    assert "扫描元数据持续秒数：13" in scope_text
    assert "未展开的覆盖内部元数据：1897 项" in scope_text
    assert "INTERNAL_SECRET" not in scope_text
    assert "private.py" not in scope_text
    assert len(scope_text) < 2_000

    html = export_to_html(task, [], "", 72, load_builtin_template("detailed"))
    word = export_to_word(task, [], "", 72)
    pdf_paragraphs: list[str] = []
    original_paragraph = report_pdf_exporter.Paragraph

    def record_pdf_paragraph(text, style, *args, **kwargs):
        pdf_paragraphs.append(str(text))
        return original_paragraph(text, style, *args, **kwargs)

    monkeypatch.setattr(report_pdf_exporter, "Paragraph", record_pdf_paragraph)
    assert export_to_pdf(task, [], "", 72).startswith(b"%PDF")
    for rendered in (
        html,
        "\n".join(paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs),
        "\n".join(pdf_paragraphs),
    ):
        assert "未展开的覆盖内部元数据" in rendered
        assert "INTERNAL_SECRET" not in rendered
        assert "private.py" not in rendered


def test_report_word_export_uses_available_chinese_font_for_styles_and_runs():
    """服务报告 Word 的样式与正文 run 都应设置可用的东亚字体。"""
    document = Document(io.BytesIO(export_to_word(_make_task(), [], "中文摘要", 72)))
    normal_style = document.styles["Normal"]
    assert normal_style.font.name == "Noto Sans SC"
    assert normal_style.element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Sans SC"

    with zipfile.ZipFile(io.BytesIO(export_to_word(_make_task(), [], "中文摘要", 72))) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:eastAsia="Noto Sans SC"' in document_xml


def test_report_word_export_keeps_code_label_with_code_block():
    """修复代码的小标题不可孤立在分页底部。"""
    issue = _make_issue(fixed_code="cursor.execute(sql, params)")
    document = Document(io.BytesIO(export_to_word(_make_task(), [issue], "摘要", 72)))
    code_label = next(paragraph for paragraph in document.paragraphs if paragraph.text == "修复代码:")
    assert code_label.paragraph_format.keep_with_next is True


def test_word_numbered_remediation_lists_restart_for_each_issue():
    """每条问题各自的编号修复步骤应从 1 开始。"""
    issues = [
        _make_issue(id=index, title=f"问题 {index}", severity="高", remediation="1. 第一步\n2. 第二步")
        for index in (1, 2)
    ]
    document = Document(io.BytesIO(export_to_word(_make_task(), issues, "摘要", 72)))
    numbered_steps = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text in {"1. 第一步", "2. 第二步"}
    ]

    assert numbered_steps == ["1. 第一步", "2. 第二步", "1. 第一步", "2. 第二步"]


def test_standard_exports_render_markdown_summary_without_showing_syntax(monkeypatch):
    summary = (
        "## 总体结论\n\n"
        "**发现两个高危问题**\n\n"
        "- 影响：可能造成越权\n"
        "- 修复：校验授权\n\n"
        "```python\n*literal = 1\n```"
    )
    html = export_to_html(_make_task(), [], summary, 72, load_builtin_template("detailed"))
    summary_box = html.split('<div class="summary-box">', 1)[1].split("</div>", 1)[0]
    assert "<h2>总体结论</h2>" in summary_box
    assert "<strong>发现两个高危问题</strong>" in summary_box
    assert "##" not in summary_box
    assert "**" not in summary_box
    assert "<li>影响：可能造成越权</li>" in summary_box
    assert "*literal = 1" in summary_box  # 代码符号属于原文，不能被 Markdown 解析吞掉。

    captured_elements = []

    class CapturedPdfDocument:
        def __init__(self, buffer, **_kwargs):
            self.buffer = buffer

        def build(self, elements):
            captured_elements.extend(elements)
            self.buffer.write(b"%PDF-test")

    monkeypatch.setattr(report_pdf_exporter, "SimpleDocTemplate", CapturedPdfDocument)
    assert export_to_pdf(_make_task(), [], summary, 72).startswith(b"%PDF")
    pdf_visible_text = "\n".join(
        element.getPlainText() for element in captured_elements if hasattr(element, "getPlainText")
    ).replace("\xa0", " ")
    assert "总体结论" in pdf_visible_text
    assert "发现两个高危问题" in pdf_visible_text
    assert "影响：可能造成越权" in pdf_visible_text
    assert "##" not in pdf_visible_text
    assert "**" not in pdf_visible_text
    assert "*literal = 1" in pdf_visible_text

    word = export_to_word(_make_task(), [], summary, 72)
    word_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs)
    assert "总体结论" in word_text
    assert "发现两个高危问题" in word_text
    assert "影响：可能造成越权" in word_text
    assert "##" not in word_text
    assert "**" not in word_text
    assert "*literal = 1" in word_text


def test_legacy_report_export_marks_missing_coverage_unknown():
    import json

    payload = json.loads(export_to_json(
        _make_task(status="failed", coverage={}), [], "扫描失败", 0,
    ))
    assert payload["task_info"]["coverage"]["coverage_status"] == "unknown"
    assert "无法确认实际扫描范围" in payload["task_info"]["coverage"]["unscanned"][0]


@pytest.mark.parametrize("coverage_status", ["partial", "failed", "unknown"])
def test_report_exports_score_reviewed_scope_when_coverage_is_not_complete(coverage_status):
    task = _make_task(
        score=100,
        coverage={"coverage_status": coverage_status, "score_available": True},
    )

    result = export_to_dict(task, [], "仍保留已取得的报告内容", 100)

    # 覆盖不完整不再隐藏分数：按已审范围规则出分并标注评分范围。
    assert result["score"] == 100
    assert result["task_info"]["score"] == 100
    assert result["statistics"]["score_breakdown"]["score_scope"] == "partial"
    assert result["statistics"]["risk_level"] != "未评估"


def test_report_export_without_any_reviewed_scope_is_not_scored():
    task = _make_task(
        status="failed", score=0, processed_files=0, total_issues=0,
        coverage={"coverage_status": "failed", "score_available": False},
    )

    result = export_to_dict(task, [], "没有任何文件完成审查", 0)

    assert result["score"] is None
    assert result["statistics"]["score_breakdown"]["score_source"] == "unavailable"
    assert result["statistics"]["risk_level"] == "未评估"


def test_report_with_ungraded_issue_does_not_receive_unearned_100_score():
    result = export_to_dict(
        _make_task(score=100),
        [_make_issue(severity="未分级")],
        "问题严重度尚未确定",
        100,
    )

    assert result["statistics"]["severity_count"]["其他"] == 1
    assert result["score"] is None
    assert result["task_info"]["score"] is None
    assert result["statistics"]["score_breakdown"]["score_source"] == "unavailable"
    assert result["statistics"]["score_breakdown"]["score_unavailable_reason"] == "ungraded_findings"
    assert "存在未分级问题" in "\n".join(result["scope_lines"])


def test_refuted_issue_is_excluded_from_score_but_remains_in_report_count():
    refuted = _make_issue(
        severity="严重",
        status="refuted",
        cvss_score=9.8,
        cwe="CWE-89",
        compliance_mapping={"owasp_top_10": ["A03:2021"]},
    )
    result = export_to_dict(
        _make_task(score=45),
        [refuted],
        "人工复核已排除该问题",
        45,
    )

    assert result["statistics"]["total_issues"] == 1
    assert result["statistics"]["severity_count"]["严重"] == 0
    assert result["score"] == 100
    assert result["statistics"]["score_breakdown"]["counts"]["严重"] == 0
    assert result["statistics"]["active_issue_count"] == 0
    assert result["statistics"]["refuted_issue_count"] == 1
    assert result["statistics"]["type_count"] == {}
    assert result["statistics"]["cwe_count"] == {}
    assert result["statistics"]["top_vulnerabilities"] == []
    assert all(
        standard["total_findings"] == 0
        for standard in result["statistics"]["compliance_summary"].values()
    )
    assert "已证伪项保留在明细中，不纳入风险汇总与评分。" in "\n".join(result["scope_lines"])


def test_refuted_finding_status_is_readable_in_html_word_and_pdf(monkeypatch):
    issue = _make_issue(severity="严重", status="refuted")
    task = _make_task()
    html = export_to_html(task, [issue], "人工复核已排除", 100, load_builtin_template("simple"))
    word = export_to_word(task, [issue], "人工复核已排除", 100)
    pdf_paragraphs: list[str] = []
    original_paragraph = report_pdf_exporter.Paragraph

    def record_pdf_paragraph(text, style, *args, **kwargs):
        pdf_paragraphs.append(str(text))
        return original_paragraph(text, style, *args, **kwargs)

    monkeypatch.setattr(report_pdf_exporter, "Paragraph", record_pdf_paragraph)
    assert export_to_pdf(task, [issue], "人工复核已排除", 100).startswith(b"%PDF")
    word_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs)
    pdf_text = "\n".join(pdf_paragraphs)

    assert "已证伪" in html
    assert "已证伪" in word_text
    assert "已证伪" in pdf_text


def test_inconsistent_coverage_and_verification_statuses_remain_separate_and_scope_score():
    result = export_to_dict(
        _make_task(coverage={
            "coverage_status": "partial",
            "verification_status": "complete",
            "score_available": True,
        }),
        [],
        "范围只完成一部分",
        100,
    )

    assert result["task_info"]["coverage"]["coverage_status"] == "partial"
    assert result["task_info"]["coverage"]["verification_status"] == "complete"
    assert result["score"] == 100
    assert result["statistics"]["score_breakdown"]["score_scope"] == "partial"
    assert "覆盖状态：部分覆盖" in "\n".join(result["scope_lines"])
    assert "验证状态：已核验通过" in "\n".join(result["scope_lines"])


def test_pdf_and_word_score_formatters_show_unavailable_as_unassessed():
    assert report_pdf_exporter._format_score(None) == "未评估"
    assert format_word_score(None) == "未评估"


def test_report_without_coverage_snapshot_scores_as_historical_task():
    result = export_to_dict(_make_task(score=88, coverage={}), [], "历史报告", 88)

    assert result["task_info"]["coverage"]["coverage_status"] == "unknown"
    assert result["score"] == 88
    assert result["task_info"]["score"] == 88
    assert result["statistics"]["score_breakdown"]["score_scope"] == "historical"


def test_export_to_json_preserves_chinese_chars():
    """JSON 应使用 ensure_ascii=False,中文不转义。"""
    raw = export_to_json(_make_task(), [_make_issue(severity="严重")], "摘要", 70)
    assert "严重" in raw
    assert "\\u" not in raw


def test_export_to_dict_returns_dict_with_required_keys():
    """dict 导出应返回字典且包含 5 个顶层键。"""
    result = export_to_dict(_make_task(), [], None, 60)
    assert isinstance(result, dict)
    assert "task_info" in result
    assert "summary" in result
    assert "score" in result
    assert "issues" in result
    assert "statistics" in result


def test_export_to_dict_keeps_structured_decompilation_evidence():
    evidence = {
        "decompilation": {
            "status": "succeeded",
            "tool": "jadx",
            "tool_version": "1.5.6",
            "input_sha256": "c" * 64,
            "output_sha256": "a" * 64,
            "exit_code": 0,
            "artifact_refs": ["decompilation-manifest"],
        }
    }
    result = export_to_dict(_make_task(), [], "摘要", 80, evidence)
    assert result["evidence"]["decompilation"]["input_sha256"] == "c" * 64
    assert result["evidence"]["decompilation"]["artifact_refs"] == ["decompilation-manifest"]


def test_export_to_dict_statistics_contains_required_fields():
    """statistics 应包含 severity_count / type_count / cwe_count / compliance_summary / top_vulnerabilities。"""
    result = export_to_dict(_make_task(), _make_mixed_issues(), "摘要", 70)
    stats = result["statistics"]
    for key in ("total_issues", "fixed_count", "severity_count", "type_count",
                "cwe_count", "compliance_summary", "top_vulnerabilities",
                "score_version", "score_breakdown", "risk_level"):
        assert key in stats, f"statistics 缺少字段 {key}"


def test_nonempty_report_uses_final_issues_as_only_score_source():
    """非空报告必须忽略任务快照与传入旧分数，按最终问题集重算。"""
    task = _make_task(
        score=5,
        total_issues=99,
        severe_issues=0,
        high_issues=99,
        score_version="legacy-v0",
        score_breakdown={"score": 5, "risk_level": "低风险"},
    )
    issues = [
        _make_issue(id=1, severity="严重"),
        _make_issue(id=2, severity="高"),
        _make_issue(id=3, severity="中"),
        _make_issue(id=4, severity="低"),
    ]

    result = export_to_dict(task, issues, "摘要", 1)

    assert result["score"] == 73
    assert result["task_info"]["score"] == 73
    assert result["task_info"]["total_issues"] == 4
    assert result["statistics"]["severity_count"] == {
        "严重": 1,
        "高": 1,
        "中": 1,
        "低": 1,
        "其他": 0,
    }
    assert result["statistics"]["score_breakdown"] == {
        "version": "severity-deduction-v1",
        "base_score": 100,
        "weights": {"严重": 15, "高": 8, "中": 3, "低": 1},
        "counts": {"严重": 1, "高": 1, "中": 1, "低": 1},
        "deductions": {"严重": 15, "高": 8, "中": 3, "低": 1},
        "total_deduction": 27,
        "raw_score": 73,
        "score": 73,
        "risk_level": "中风险",
        "score_source": "review_issues",
        "score_scope": "complete",
    }


@pytest.mark.parametrize("explicit_score", [100, 72])
def test_empty_report_preserves_explicit_historical_score(explicit_score):
    """空问题历史报告保留显式评分，同时保持统一权重与风险说明。"""
    result = export_to_dict(
        _make_task(score=explicit_score, total_issues=0),
        [],
        "",
        explicit_score,
    )

    breakdown = result["statistics"]["score_breakdown"]
    assert result["score"] == explicit_score
    assert result["task_info"]["score"] == explicit_score
    assert breakdown["score"] == explicit_score
    assert breakdown["weights"] == {"严重": 15, "高": 8, "中": 3, "低": 1}
    assert breakdown["score_source"] == "task_explicit_empty_report"
    assert breakdown["risk_level"] == ("低风险" if explicit_score == 100 else "中风险")


def test_report_discloses_review_candidates_and_rejections_without_changing_score():
    """原扣分口径包含未核实与驳回项，报告必须明示而非冒充已验证漏洞评分。"""
    issues = [
        _make_issue(id=1, severity="严重", human_review_status="pending"),
        _make_issue(id=2, severity="高", human_review_status="rejected", status="ignored"),
        _make_issue(id=3, severity="中", human_review_status="accepted", confirmation_count=99),
        _make_issue(id=4, severity="低", human_review_status="evidence_requested"),
        _make_issue(id=5, severity="低", human_review_status=None, evidence_quality="verified"),
    ]
    result = export_to_dict(_make_task(), issues, "", 1)

    assert result["score"] == 72  # 15 + 8 + 3 + 1 + 1，保留现有算法与问题范围。
    assert result["statistics"]["score_breakdown"]["total_deduction"] == 28
    basis = result["statistics"]["assessment_basis"]
    assert basis["score_source"] == "review_issues"
    assert basis["issue_scope"] == "reported_review_issues_excluding_refuted_from_risk_aggregates"
    assert basis["human_review_counts"] == {
        "pending": 1, "evidence_requested": 1, "accepted": 1,
        "rejected": 1, "not_required": 0, "unrecorded": 1,
    }
    assert basis["confirmation_count_semantics"] == "retained_source_count_not_vulnerability_verification"
    assert basis["evidence_quality_semantics"] == "source_quote_match_not_condition_or_impact_verification"
    assert basis["manual_status_semantics"] == "manual_disposition_not_retest_verification"
    assert basis["domestic_database_freshness"] == "not_verified_by_this_export"
    assert basis["official_rating"] == "not_provided_by_this_export"
    assert result["issues"][4]["human_review_status"] is None  # 历史未知不补成已核实。
    text = "\n".join(result["scope_lines"])
    assert "待核实 2 项、人工驳回 1 项仍纳入现有评分口径" in text
    assert "confirmation_count 来源数" in text
    assert "不代表漏洞条件或影响已验证" in text
    assert "已修复标记不等于复测通过" in text
    assert "不能据此声明已同步最新库或获得权威认证" in text


@pytest.mark.parametrize("explicit_score", [100, 72])
def test_empty_historical_score_disclosure_does_not_infer_verified_coverage(explicit_score):
    result = export_to_dict(_make_task(score=explicit_score), [], "", explicit_score)

    assert result["score"] == explicit_score
    assert result["statistics"]["assessment_basis"]["score_source"] == "task_explicit_empty_report"
    assert all(value == 0 for value in result["statistics"]["assessment_basis"]["human_review_counts"].values())
    assert "空问题报告保留历史显式分数，不代表已验证无漏洞" in "\n".join(result["scope_lines"])


@pytest.mark.parametrize("template_type", ["simple", "detailed", "compliance"])
def test_builtin_html_reports_display_evidence_and_score_scope(template_type):
    issue = _make_issue(human_review_status="pending", confirmation_count=5)
    html = export_to_html(_make_task(), [issue], "", 1, load_builtin_template(template_type))

    assert "平台内部问题扣分" in html
    assert "待核实 1 项、人工驳回 0 项仍纳入现有评分口径" in html
    assert "不代表漏洞条件或影响已验证" in html
    assert "不能据此声明已同步最新库或获得权威认证" in html
    assert "92/100" in html  # 待核实的高等级问题仍按现有权重扣 8 分。


def test_word_report_retains_assessment_scope_without_upgrading_manual_acceptance():
    word = export_to_word(
        _make_task(), [_make_issue(human_review_status="accepted", status="fixed")], "", 1,
    )
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    assert "平台内部问题扣分" in document_xml
    assert "不代表漏洞条件或影响已验证" in document_xml
    assert "已修复标记不等于复测通过" in document_xml


@pytest.mark.parametrize(
    "status,expected_key",
    [
        ("pending", "pending"), ("evidence_requested", "evidence_requested"),
        ("accepted", "accepted"), ("rejected", "rejected"),
        ("not_required", "not_required"), (None, "unrecorded"),
        ("", "unrecorded"), ("historical_unknown", "unrecorded"),
    ],
)
def test_json_assessment_metadata_preserves_recorded_and_unknown_review_statuses(status, expected_key):
    import json

    issue = _make_issue(human_review_status=status, confirmation_count=99, confidence=1.0)
    result = json.loads(export_to_json(_make_task(), [issue], "", 1))
    counts = result["statistics"]["assessment_basis"]["human_review_counts"]

    assert counts[expected_key] == 1
    assert sum(counts.values()) == 1
    assert result["score"] == 92
    assert result["issues"][0]["human_review_status"] == status
    assert result["issues"][0]["confirmation_count"] == 99
    assert result["statistics"]["assessment_basis"]["confidence_semantics"] == (
        "model_estimate_not_vulnerability_verification"
    )


def test_pdf_report_actually_builds_visible_assessment_scope(monkeypatch):
    from reportlab.platypus import SimpleDocTemplate

    original_build = SimpleDocTemplate.build
    visible_text = []

    def capture_build(doc, flowables, *args, **kwargs):
        visible_text.extend(
            element.getPlainText() for element in flowables if hasattr(element, "getPlainText")
        )
        return original_build(doc, flowables, *args, **kwargs)

    monkeypatch.setattr(SimpleDocTemplate, "build", capture_build)
    pdf = export_to_pdf(_make_task(), [_make_issue(human_review_status="rejected")], "", 1)

    assert pdf.startswith(b"%PDF")
    text = "\n".join(visible_text)
    assert "人工驳回 1 项仍纳入现有评分口径" in text
    assert "不代表漏洞条件或影响已验证" in text
    assert "不能据此声明已同步最新库或获得权威认证" in text


def test_task_name_and_name_are_bidirectionally_compatible():
    """新旧报告调用方使用 name 或 task_name 都得到两个稳定别名。"""
    result = export_to_dict(
        {"id": 1, "name": "旧字段任务", "score": 100},
        [],
        "",
        100,
    )

    assert result["task_info"]["name"] == "旧字段任务"
    assert result["task_info"]["task_name"] == "旧字段任务"


# ============ HTML 导出(3 套模板) ============

def test_export_to_html_simple_renders_with_key_content():
    """简洁版模板渲染应包含任务名、分数与问题列表。"""
    task = _make_task(task_name="简洁任务")
    issues = [_make_issue(title="问题A"), _make_issue(id=2, title="问题B", severity="中")]
    tpl = load_builtin_template("simple")
    html = export_to_html(task, issues, "总体评价", 75, tpl)
    assert "简洁任务" in html
    assert "89" in html
    assert "问题A" in html
    assert "问题B" in html
    assert "<table" in html


def test_export_to_html_detailed_renders_with_key_content():
    """详细版模板渲染应包含任务名、分数、问题描述与修复代码。"""
    task = _make_task(task_name="详细任务")
    issues = [_make_issue(title="SQL注入")]
    tpl = load_builtin_template("detailed")
    html = export_to_html(task, issues, "详细评价", 65, tpl)
    assert "详细任务" in html
    assert "SQL注入" in html
    assert "参数化查询" in html  # suggestion
    assert "<pre>" in html and "<code>" in html  # 修复代码
    assert "SQL注入" in html  # 标题(已覆盖)


def test_all_html_templates_render_canonical_score_breakdown():
    """三套 HTML 报告都必须可读地解释 15/8/3/1 评分。"""
    issues = [
        _make_issue(id=1, severity="严重"),
        _make_issue(id=2, severity="高"),
    ]

    for template_type in ("simple", "detailed", "compliance"):
        html = export_to_html(
            _make_task(score=1),
            issues,
            "",
            1,
            load_builtin_template(template_type),
        )
        assert "评分明细" in html
        assert "严重 × 15" in html
        assert "高 × 8" in html
        assert "总扣分" in html
        assert "77/100" in html


def test_export_to_html_supports_legacy_template_context():
    """历史模板的 task/metrics/compliance_summary 变量应继续可用。"""

    task = _make_task(total_files=4)
    issues = [
        _make_issue(id=1, severity="严重"),
        _make_issue(id=2, severity="high"),
        _make_issue(id=3, severity="信息"),
    ]
    template = (
        "{{ task.task_name }}|{{ metrics.total_files }}|"
        "{{ metrics.total_issues }}|{{ metrics.severity_counts.critical }}|"
        "{{ metrics.severity_counts.high }}|{{ metrics.severity_counts.info }}|"
        "{{ compliance_summary.iso27001.total_findings }}"
    )

    rendered = export_to_html(task, issues, "旧模板", 66, template)

    assert rendered == "示例审查任务|4|3|1|1|1|3"


def test_export_to_html_blocks_unsafe_template_attribute_access():
    """数据库模板只能访问报告上下文，不能遍历 Python 运行时对象。"""
    with pytest.raises(SecurityError):
        export_to_html(
            _make_task(task_name="sandbox"),
            [],
            "summary",
            100,
            "{{ task_info.__class__.__mro__ }}",
        )


def test_export_to_html_compliance_renders_with_key_content():
    """合规版模板渲染应包含任务名、4 标准概览与合规条款编号。"""
    task = _make_task(project_name="合规项目")
    issues = [_make_issue(cwe="CWE-89"), _make_issue(id=2, cwe="CWE-79", title="XSS")]
    tpl = load_builtin_template("compliance")
    html = export_to_html(task, issues, "合规评价", 55, tpl)
    assert "合规项目" in html
    assert "ISO 27001" in html
    assert "GDPR" in html
    assert "PCI-DSS" in html
    assert "HIPAA" in html
    assert "A.8.25" in html or "A.8.28" in html  # 合规条款编号
    assert "<code>" in html


def test_export_to_html_simple_uses_inline_css():
    """简洁版模板应使用内联 CSS(<style> 标签),不依赖外部样式表。"""
    tpl = load_builtin_template("simple")
    html = export_to_html(_make_task(), [], "", 80, tpl)
    assert "<style>" in html
    assert "<link" not in html  # 无外部样式表


def test_export_to_html_detailed_contains_evidence_and_exploit():
    """详细版应渲染 evidence 与 exploit_scenario。"""
    issues = [_make_issue(evidence="evidence_line", exploit_scenario="攻击场景X")]
    tpl = load_builtin_template("detailed")
    html = export_to_html(_make_task(), issues, "", 70, tpl)
    assert "evidence_line" in html
    assert "攻击场景X" in html


def test_export_to_html_detailed_includes_anchor_links():
    """详细版应为每个问题生成锚点 id。"""
    issues = [_make_issue(id=101), _make_issue(id=102, title="第二个")]
    tpl = load_builtin_template("detailed")
    html = export_to_html(_make_task(), issues, "", 70, tpl)
    assert 'id="issue-101"' in html
    assert 'id="issue-102"' in html
    assert "#issue-101" in html  # TOC 链接


def test_export_to_html_compliance_contains_cvss_distribution_and_top10():
    """合规版应包含 CVSS 分布章节与 Top 10 漏洞表格。"""
    issues = _make_mixed_issues()
    tpl = load_builtin_template("compliance")
    html = export_to_html(_make_task(), issues, "", 50, tpl)
    assert "CVSS 评分分布" in html
    assert "Top 10" in html
    assert "cvss-badge" in html or "cvss-bar" in html


def test_export_to_html_compliance_renders_cvss_vector():
    """合规版应渲染 cvss_vector 字段。"""
    issues = [_make_issue(cvss_vector="AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H")]
    tpl = load_builtin_template("compliance")
    html = export_to_html(_make_task(), issues, "", 50, tpl)
    assert "AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H" in html


def test_template_rendering_includes_cvss_score_field():
    """详细版与合规版均应包含 cvss_score 数值。"""
    issues = [_make_issue(cvss_score=8.6)]
    for tpl_type in ("detailed", "compliance"):
        tpl = load_builtin_template(tpl_type)
        html = export_to_html(_make_task(), issues, "", 60, tpl)
        assert "8.6" in html


def test_template_rendering_includes_compliance_mapping():
    """合规版应渲染 compliance_mapping 中的条款编号。"""
    issues = [_make_issue(compliance_mapping={
        "iso27001": ["A.8.25"], "gdpr": ["Art.32"],
        "pci_dss": ["Req-6.4.1"], "hipaa": ["§164.312(a)"],
    })]
    tpl = load_builtin_template("compliance")
    html = export_to_html(_make_task(), issues, "", 60, tpl)
    assert "Art.32" in html
    assert "Req-6.4.1" in html


# ============ 按严重度分组 ============

def test_group_issues_by_severity_returns_four_groups_in_order():
    """按严重度分组应返回 严重/高/中/低 4 组,顺序固定。"""
    issues = [
        _make_issue(id=1, severity="低"),
        _make_issue(id=2, severity="严重"),
        _make_issue(id=3, severity="中"),
        _make_issue(id=4, severity="高"),
    ]
    grouped = _group_issues_by_severity(issues)
    assert list(grouped.keys())[:4] == ["严重", "高", "中", "低"]
    assert len(grouped["严重"]) == 1
    assert len(grouped["高"]) == 1
    assert len(grouped["中"]) == 1
    assert len(grouped["低"]) == 1
    assert grouped["严重"][0]["id"] == 2


def test_group_issues_by_severity_handles_unknown_severity():
    """未知严重度应归入 '其他' 组。"""
    issues = [_make_issue(id=1, severity="未知"), _make_issue(id=2, severity="严重")]
    grouped = _group_issues_by_severity(issues)
    assert "其他" in grouped
    assert len(grouped["其他"]) == 1
    assert grouped["其他"][0]["id"] == 1


def test_group_issues_by_severity_empty_input():
    """空列表分组应返回空分组结构。"""
    grouped = _group_issues_by_severity([])
    assert grouped["严重"] == []
    assert grouped["高"] == []
    assert grouped["中"] == []
    assert grouped["低"] == []


# ============ 按类型分组 ============

def test_group_issues_by_type_returns_nine_types():
    """按类型分组应返回 9 种 issue_type。"""
    issues = _make_mixed_issues()
    grouped = _group_issues_by_type(issues)
    assert len(grouped) == 9
    expected_types = {"安全漏洞", "潜在Bug", "性能问题", "代码风格", "可维护性",
                      "兼容性", "并发问题", "资源泄漏", "配置问题"}
    assert set(grouped.keys()) == expected_types


def test_group_issues_by_type_preserves_issues_count():
    """每个类型分组应保留正确的 issue 数量。"""
    issues = [
        _make_issue(id=1, issue_type="安全漏洞"),
        _make_issue(id=2, issue_type="安全漏洞"),
        _make_issue(id=3, issue_type="性能问题"),
    ]
    grouped = _group_issues_by_type(issues)
    assert len(grouped["安全漏洞"]) == 2
    assert len(grouped["性能问题"]) == 1


def test_group_issues_by_type_sorted_by_key():
    """类型分组键应按字典序升序排列。"""
    issues = [
        _make_issue(id=1, issue_type="性能问题"),
        _make_issue(id=2, issue_type="安全漏洞"),
        _make_issue(id=3, issue_type="代码风格"),
    ]
    grouped = _group_issues_by_type(issues)
    keys = list(grouped.keys())
    assert keys == sorted(keys)


# ============ 合规摘要聚合 ============

def test_build_compliance_summary_aggregates_four_standards():
    """合规摘要应聚合 4 套标准,且命中条目数正确。"""
    issues = [
        _make_issue(id=1, cwe="CWE-89", compliance_mapping=None),
        _make_issue(id=2, cwe="CWE-79", compliance_mapping=None),
    ]
    summary = _build_compliance_summary(issues)
    assert set(summary.keys()) == {"iso27001", "gdpr", "pci_dss", "hipaa"}
    # 两条 issue 均映射到 4 标准
    assert summary["iso27001"]["total_findings"] == 2
    assert summary["gdpr"]["total_findings"] == 2
    assert summary["pci_dss"]["total_findings"] == 2
    assert summary["hipaa"]["total_findings"] == 2
    # covered 字段非空
    assert len(summary["iso27001"]["covered_controls"]) > 0
    assert len(summary["gdpr"]["covered_articles"]) > 0
    assert len(summary["pci_dss"]["covered_requirements"]) > 0
    assert len(summary["hipaa"]["covered_sections"]) > 0


def test_build_compliance_summary_uses_explicit_mapping_when_provided():
    """当 issue 提供 compliance_mapping 时应直接使用,不按 cwe 补全。"""
    issues = [_make_issue(
        cwe="CWE-89",
        compliance_mapping={
            "iso27001": ["A.8.25"], "gdpr": [], "pci_dss": [], "hipaa": [],
        },
    )]
    summary = _build_compliance_summary(issues)
    assert summary["iso27001"]["total_findings"] == 1
    assert summary["gdpr"]["total_findings"] == 0
    assert "A.8.25" in summary["iso27001"]["covered_controls"]


def test_build_compliance_summary_empty_issues():
    """空问题列表时合规摘要所有 total_findings 为 0,covered 为空。"""
    summary = _build_compliance_summary([])
    for std in ("iso27001", "gdpr", "pci_dss", "hipaa"):
        assert summary[std]["total_findings"] == 0


def test_build_compliance_summary_deduplicates_codes():
    """合规条款编号应去重。"""
    issues = [
        _make_issue(id=1, cwe="CWE-89", compliance_mapping=None),
        _make_issue(id=2, cwe="CWE-89", compliance_mapping=None),
    ]
    summary = _build_compliance_summary(issues)
    # CWE-89 的 iso27001 映射为 ["A.8.25", "A.8.26", "A.8.28", "A.8.29"],去重后仍为 4 条
    assert len(summary["iso27001"]["covered_controls"]) == 4
    assert summary["iso27001"]["total_findings"] == 2


# ============ Top N 漏洞 ============

def test_build_top_vulnerabilities_sorted_by_cvss_desc():
    """Top N 应按 cvss_score 降序排序。"""
    issues = [
        _make_issue(id=1, cvss_score=5.0),
        _make_issue(id=2, cvss_score=9.8),
        _make_issue(id=3, cvss_score=7.2),
    ]
    top = _build_top_vulnerabilities(issues, top_n=3)
    assert [t["id"] for t in top] == [2, 3, 1]
    assert top[0]["cvss_score"] == 9.8


def test_build_top_vulnerabilities_respects_top_n():
    """Top N 应截断到指定数量。"""
    issues = _make_mixed_issues()
    top = _build_top_vulnerabilities(issues, top_n=3)
    assert len(top) == 3


def test_build_top_vulnerabilities_handles_none_cvss():
    """cvss_score 为 None 的 issue 应排到最后。"""
    issues = [
        _make_issue(id=1, cvss_score=None),
        _make_issue(id=2, cvss_score=4.0),
    ]
    top = _build_top_vulnerabilities(issues, top_n=2)
    assert top[0]["id"] == 2
    assert top[1]["id"] == 1


def test_build_top_vulnerabilities_empty_returns_empty_list():
    """空列表应返回空列表。"""
    assert _build_top_vulnerabilities([], top_n=10) == []


def test_build_top_vulnerabilities_zero_top_n_returns_empty():
    """top_n<=0 应返回空列表。"""
    issues = [_make_issue(cvss_score=9.0)]
    assert _build_top_vulnerabilities(issues, top_n=0) == []


# ============ 空场景 / 单问题 / 混合场景 ============

def test_empty_issues_does_not_raise_and_statistics_zero():
    """空问题列表场景下导出函数不应报错,统计字段为 0。"""
    task = _make_task()
    for fn in (export_to_json, export_to_dict):
        fn(task, [], None, 100)
    # HTML 也应能渲染空场景
    tpl = load_builtin_template("simple")
    html = export_to_html(task, [], "", 100, tpl)
    assert "未发现问题" in html

    payload = export_to_dict(task, [], None, 100)
    assert payload["statistics"]["total_issues"] == 0
    assert payload["statistics"]["fixed_count"] == 0
    assert payload["statistics"]["severity_count"]["严重"] == 0
    assert payload["statistics"]["severity_count"]["高"] == 0
    assert payload["issues"] == []
    assert payload["statistics"]["top_vulnerabilities"] == []


def test_single_issue_scenario():
    """单个问题场景应正确导出。"""
    issues = [_make_issue(id=7, title="单点问题", severity="严重")]
    result = export_to_dict(_make_task(), issues, "摘要", 40)
    assert result["statistics"]["total_issues"] == 1
    assert result["statistics"]["severity_count"]["严重"] == 1
    assert len(result["issues"]) == 1
    assert result["issues"][0]["id"] == 7


def test_multi_issue_mixed_scenario_sorts_by_severity():
    """多问题混合场景应按严重度排序(严重>高>中>低)。"""
    issues = [
        _make_issue(id=1, severity="低"),
        _make_issue(id=2, severity="严重"),
        _make_issue(id=3, severity="中"),
        _make_issue(id=4, severity="高"),
    ]
    result = export_to_dict(_make_task(), issues, "摘要", 50)
    severities = [it["severity"] for it in result["issues"]]
    assert severities == ["严重", "高", "中", "低"]


def test_build_report_context_includes_top_vulnerabilities_capped_at_10():
    """上下文中 top_vulnerabilities 最多 10 条。"""
    issues = [_make_issue(id=i, cvss_score=float(10 - i)) for i in range(1, 15)]
    ctx = _build_report_context(_make_task(), issues, "", 50)
    assert len(ctx["statistics"]["top_vulnerabilities"]) == 10
    # 第一条 cvss 最高
    assert ctx["statistics"]["top_vulnerabilities"][0]["cvss_score"] == 9.0


# ============ Pydantic 模型输入兼容性 ============

def test_export_to_html_with_issueout_pydantic_model():
    """export_to_html 应兼容 IssueOut Pydantic 模型输入。"""
    issue_dict = _make_issue()
    issue_out = IssueOut(**issue_dict)
    tpl = load_builtin_template("detailed")
    html = export_to_html(_make_task(), [issue_out], "摘要", 70, tpl)
    assert "SQL 注入漏洞" in html
    assert "参数化查询" in html


def test_export_to_dict_with_issueout_normalizes_fields():
    """dict 导出应正确归一化 IssueOut 模型字段。"""
    issue_out = IssueOut(**_make_issue())
    result = export_to_dict(_make_task(), [issue_out], None, 70)
    assert result["issues"][0]["cwe"] == "CWE-89"
    assert result["issues"][0]["cvss_score"] == 8.6
    assert isinstance(result["issues"][0]["compliance_mapping"], dict)


# ============ 内置模板载入 ============

def test_load_builtin_template_returns_content_for_each_type():
    """3 套内置模板均能载入且非空。"""
    for tpl_type in ("simple", "detailed", "compliance"):
        content = load_builtin_template(tpl_type)
        assert isinstance(content, str)
        assert len(content) > 0
        assert "<!DOCTYPE html>" in content


def test_load_builtin_template_invalid_type_raises():
    """不支持的模板类型应抛出 FileNotFoundError。"""
    with pytest.raises(FileNotFoundError):
        load_builtin_template("unknown_type")


# ============ JSON 序列化与 datetime ============

def test_json_serialization_handles_datetime():
    """JSON 导出应正确序列化 datetime 对象。"""
    task = _make_task(create_time=datetime(2026, 6, 25, 10, 0, 0))
    issues = [_make_issue(create_time=datetime(2026, 6, 25, 10, 5, 0))]
    raw = export_to_json(task, issues, "摘要", 70)
    import json
    payload = json.loads(raw)
    assert "2026-06-25T10:00:00" in payload["task_info"]["create_time"]
    assert "2026-06-25T10:05:00" in payload["issues"][0]["create_time"]


def test_export_to_dict_normalizes_references_json_to_list():
    """references_json 缺失时应归一化为空列表,避免模板渲染错误。"""
    issue = _make_issue(references_json=None)
    issue.pop("references_json", None)
    result = export_to_dict(_make_task(), [issue], None, 70)
    assert result["issues"][0]["references_json"] == []


def test_export_to_dict_normalizes_compliance_mapping_to_dict():
    """compliance_mapping 缺失时应归一化为空字典。"""
    issue = _make_issue()
    issue.pop("compliance_mapping", None)
    result = export_to_dict(_make_task(), [issue], None, 70)
    assert result["issues"][0]["compliance_mapping"] == {}


def test_export_to_html_all_three_templates_with_empty_issues():
    """3 套模板在空问题列表下均能成功渲染不报错。"""
    task = _make_task()
    for tpl_type in ("simple", "detailed", "compliance"):
        tpl = load_builtin_template(tpl_type)
        html = export_to_html(task, [], "", 90, tpl)
        assert "<!DOCTYPE html>" in html
        assert "未发现问题" in html or "暂无高危漏洞" in html


def _decompilation_evidence() -> Dict[str, Any]:
    return {
        "decompilation": {
            "status": "succeeded",
            "tool": "jadx",
            "tool_version": "1.5.6",
            "input_sha256": "98f7476e" + "1" * 56,
            "input_artifact_sha256s": ["3a47fa04" + "3" * 56],
            "output_sha256": "8e9350b2" + "2" * 56,
            "output_file_count": 2,
            "output_size_bytes": 1369,
            "exit_code": 0,
            "log_ref": "worker.log",
            "artifact_refs": ["decompilation-manifest"],
        }
    }


def test_html_templates_render_structured_decompilation_evidence():
    """三套 HTML 报告必须实际展示反编译哈希和工具版本。"""
    for template_type in ("simple", "detailed", "compliance"):
        html = export_to_html(
            _make_task(), [], "摘要", 90,
            load_builtin_template(template_type),
            _decompilation_evidence(),
        )
        assert "反编译证据" in html
        assert "jadx 1.5.6" in html
        assert "98f7476e" in html
        assert "3a47fa04" in html
        assert "8e9350b2" in html


def test_pdf_and_word_render_structured_decompilation_evidence():
    """PDF/Word 报告必须包含可人工核对的反编译哈希。"""
    evidence = _decompilation_evidence()
    pdf = export_to_pdf(_make_task(), [], "摘要", 90, evidence=evidence)
    assert b"98f7476e" in pdf
    assert b"3a47fa04" in pdf
    assert b"8e9350b2" in pdf

    word = export_to_word(_make_task(), [], "摘要", 90, evidence=evidence)
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "反编译证据" in document_xml
    assert "98f7476e" in document_xml
    assert "3a47fa04" in document_xml
    assert "8e9350b2" in document_xml


def test_pdf_and_word_render_full_issue_audit_metadata_and_score_breakdown():
    """可下载报告必须保留复核漏洞所需的全部可读事实。"""
    issue = _make_issue(
        evidence="EVIDENCE_TOKEN_20260827",
        exploit_scenario="EXPLOIT_TOKEN_20260827",
        suggestion="SUGGESTION_TOKEN_20260827",
        remediation="REMEDIATION_TOKEN_20260827",
        cvss_vector="AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
        cvss_version="3.1",
        cvss_source="vector",
        source="llm+static",
        confirmation_count=2,
        finding_fingerprint="FINGERPRINT_TOKEN_20260827",
        status="fixed",
        handled_by=7,
        handled_at="2026-08-27T01:02:03",
        references_json=["https://example.com/REFERENCE_TOKEN_20260827"],
    )

    pdf = export_to_pdf(_make_task(score=1), [issue], "摘要", 1)
    assert pdf.startswith(b"%PDF")

    # 检查进入 ReportLab 文档树的可见文本，避免依赖本机 pdftotext 等外部程序。
    font_name = _ensure_chinese_font()
    issue_text = "\n".join(
        element.getPlainText()
        for element in _build_issue_section(
            issue,
            1,
            _build_styles(font_name),
            font_name,
        )
        if hasattr(element, "getPlainText")
    )
    for token in (
        "EVIDENCE_TOKEN_20260827",
        "EXPLOIT_TOKEN_20260827",
        "SUGGESTION_TOKEN_20260827",
        "REMEDIATION_TOKEN_20260827",
        "FINGERPRINT_TOKEN_20260827",
        "2026-08-27T01:02:03",
        "REFERENCE_TOKEN_20260827",
        "评分向量：AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
        "向量计算",
        "模型 + 静态规则",
        "已修复",
    ):
        assert token in issue_text

    word = export_to_word(_make_task(score=1), [issue], "摘要", 1)
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    for token in (
        "评分明细",
        "severity-deduction-v1",
        "严重 × 15",
        "EVIDENCE_TOKEN_20260827",
        "EXPLOIT_TOKEN_20260827",
        "SUGGESTION_TOKEN_20260827",
        "REMEDIATION_TOKEN_20260827",
        "评分向量：AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
        "向量计算",
        "模型 + 静态规则",
        "独立确认数: 2",
        "FINGERPRINT_TOKEN_20260827",
        "已修复",
        "处理人: 7",
        "2026-08-27T01:02:03",
        "REFERENCE_TOKEN_20260827",
    ):
        assert token in document_xml


@pytest.mark.parametrize("template_type", ["simple", "detailed", "compliance"])
def test_failed_report_without_score_never_claims_no_issues(template_type: str) -> None:
    """Codex 4.1.16 P1：失败且无有效分数的报告不得写“本次审查未发现问题”。"""
    task = _make_task(
        status="failed", processed_files=0, total_issues=0, score=None,
        coverage={"coverage_status": "failed"},
    )
    html = export_to_html(task, [], "", None, load_builtin_template(template_type))
    assert "本次审查未发现问题" not in html
    assert "不能据此判定无问题" in html
    for blob in (export_to_pdf(task, [], "", None), export_to_word(task, [], "", None)):
        assert blob


def test_empty_issue_text_distinguishes_clean_partial_and_inconclusive() -> None:
    clean = _build_report_context(_make_task(total_issues=0), [], "", 100)
    assert clean["empty_issues_text"] == "本次审查未发现问题。"
    partial = _build_report_context(_make_task(
        status="failed", processed_files=1, total_issues=0,
        coverage={"coverage_status": "failed", "completed_files": 1},
    ), [], "", None)
    assert "已审范围内未发现问题" in partial["empty_issues_text"]
    failed = _build_report_context(_make_task(
        status="failed", processed_files=0, total_issues=0, coverage={"coverage_status": "failed"},
    ), [], "", None)
    assert "不能据此判定无问题" in failed["empty_issues_text"]


def test_unscored_export_states_the_concrete_reason_in_scope_lines():
    """不出分的导出正文必须写出具体原因，而不是只显示“未评估”。"""
    task = _make_task(status="failed", processed_files=0, total_issues=0, severe_issues=0, high_issues=0,
                      medium_issues=0, low_issues=0, score=None, coverage={"coverage_status": "partial"})
    result = export_to_dict(task, [], "", None)

    assert result["score"] is None
    assert result["statistics"]["score_breakdown"]["score_unavailable_detail"]
    joined = "\n".join(result["scope_lines"])
    assert "不出分原因：" in joined
    assert "没有任何已完成审查的文件或有效发言" in joined


def test_domain_score_section_states_scope_or_reason():
    from app.services.domain_report_exporter import insert_section_after_task_info, score_section_lines

    scored = score_section_lines(80, available=True, scope_label="仅覆盖已测范围", rule="R", finding_count=2,
                                 severity_counts={"高": 1, "低": 1})
    assert "评分：80 / 100" in scored and "评分范围：仅覆盖已测范围" in scored
    assert any("记录 2 项发现（高 1 项、低 1 项）" in line for line in scored)
    unscored = score_section_lines(None, available=False, reason="安全配置探测阶段未完成", finding_count=0)
    assert "评分：未评估" in unscored and "不出分原因：安全配置探测阶段未完成" in unscored
    assert any("不等于已证明不存在风险" in line for line in unscored)
    markdown = insert_section_after_task_info("# 标题\n\n## 任务信息\n任务：A\n\n## 正文\n内容\n", scored)
    assert markdown.index("## 任务信息") < markdown.index("## 评分与结论") < markdown.index("## 正文")

"""AI report Markdown must render as readable, escaped document formatting."""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import KeepTogether, Table

from app.services import domain_report_exporter, report_exporter
from app.services.domain_report_exporter import _coverage_lines, export_markdown_report
from app.services.report_exporter import export_to_html, load_builtin_template
from app.services.report_markdown import MarkdownSegment, markdown_to_blocks, reportlab_markup
from app.services.report_pdf_exporter import (
    _append_markdown_elements,
    _build_styles,
    _build_summary_markdown_elements,
    _ensure_chinese_font,
    _paragraph,
    export_to_pdf,
)
from app.services.report_word_exporter import export_to_word

SUMMARY = """# 总体结论

发现 **高危问题**，建议先修复 `access check`。

- 核对输入边界
- 补充回归测试

[官方记录](https://www.cnvd.org.cn/flaw/show/CNVD-2026-40686)
[无效链接](javascript:alert(1))

原始输入：<script>alert('xss')</script>
"""

TABLE_SUMMARY = """| 项目 | 说明 |
| --- | --- |
| **问题** | SQL *injection |
| 来源 | CNVD |
"""

DOMAIN_TABLE_REPORT = """## 核验结果

| 项目 | 观测值 | 判定 |
| --- | --- | --- |
| **黑盒状态** | 未执行 | 本次只请求白盒验证 |
| 超长证据 | 长证据内容用于确认 PDF 和 Word 会自动换行，而不是把表格竖线当作正文输出。 | 已记录 |
"""


def _task():
    return {
        "id": 10,
        "task_name": "Markdown 报告验证",
        "project_name": "合成项目",
        "review_type": "standard",
        "status": "success",
        "total_files": 1,
        "total_issues": 0,
        "score": 100,
        "coverage": {"coverage_status": "complete"},
    }


def test_all_builtin_html_templates_render_safe_formatted_markdown():
    for template_type in ("simple", "detailed", "compliance"):
        html = export_to_html(_task(), [], SUMMARY, 100, load_builtin_template(template_type))

        assert "<h1>总体结论</h1>" in html
        assert "<strong>高危问题</strong>" in html
        assert "<ul>" in html and "<li>核对输入边界</li>" in html
        assert '<a href="https://www.cnvd.org.cn/flaw/show/CNVD-2026-40686">官方记录</a>' in html
        assert "javascript:" not in html
        assert "<script>alert" not in html
        assert "&lt;script&gt;" in html
        assert "**高危问题**" not in html


def test_pdf_export_consumes_markdown_without_leaking_emphasis_markers():
    pdf = export_to_pdf(_task(), [], SUMMARY, 100)

    assert pdf.startswith(b"%PDF-")
    assert b"**" not in pdf
    assert b"javascript:" not in pdf


def test_pdf_mixed_language_paragraph_uses_latin_metrics_and_keeps_code_monospace():
    styles = _build_styles(_ensure_chinese_font())
    body = _paragraph("路径 backend/app.py，时间 2026-10-16T06:40:20", styles["normal"])
    code = _paragraph("agent_service.py", styles["code"])

    body_fonts = {fragment.fontName for fragment in body.frags}
    code_fonts = {fragment.fontName for fragment in code.frags}
    assert "Vera" in body_fonts
    assert "STSong-Light" in body_fonts
    assert code_fonts == {"Courier"}
    pdf_buffer = BytesIO()
    canvas = Canvas(pdf_buffer)
    body.wrapOn(canvas, 500, 100)
    code.wrapOn(canvas, 500, 100)
    body.drawOn(canvas, 48, 760)
    code.drawOn(canvas, 48, 730)
    canvas.save()
    pdf = pdf_buffer.getvalue()
    assert b"BitstreamVeraSans-Roman" in pdf
    assert b"/FontFile2" in pdf


def test_report_scope_explains_coverage_is_not_line_by_line_semantic_review():
    context = {
        "task_info": {
            "status": "success",
            "processed_files": 1,
            "total_files": 1,
            "coverage": {"coverage_status": "complete"},
        },
        "source": {"source_type": "discuss"},
        "files": [],
        "statistics": {
            "total_issues": 0,
            "refuted_issue_count": 0,
            "assessment_basis": {
                "human_review_counts": {
                    "pending": 0,
                    "evidence_requested": 0,
                    "rejected": 0,
                    "unrecorded": 0,
                },
                "score_source": "review_issues",
            },
            "score_breakdown": {},
        },
    }

    lines = report_exporter._build_scope_lines(context)

    assert any("不应单独解读为逐行语义审查完整" in line for line in lines)


def test_word_export_uses_native_markdown_formatting_and_safe_links():
    word = export_to_word(_task(), [], SUMMARY, 100)
    document = Document(BytesIO(word))

    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    assert "总体结论" in paragraph_text
    assert "发现 高危问题，建议先修复 access check。" in paragraph_text
    assert any(paragraph.style.name == "List Bullet" for paragraph in document.paragraphs)
    assert all("**" not in text for text in paragraph_text)
    bold_runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.bold]
    assert any(run.text == "高危问题" for run in bold_runs)
    external_links = [
        relation.target_ref
        for relation in document.part.rels.values()
        if relation.reltype.endswith("/hyperlink")
    ]
    assert external_links == ["https://www.cnvd.org.cn/flaw/show/CNVD-2026-40686"]


def test_summary_tables_render_as_tables_in_html_word_and_pdf_without_raw_markers():
    html = export_to_html(_task(), [], TABLE_SUMMARY, 100, load_builtin_template("detailed"))
    assert '<table class="report-summary-table">' in html
    assert "<th>项目</th>" in html and "<td>SQL injection</td>" in html
    assert "*injection" not in html

    document = Document(BytesIO(export_to_word(_task(), [], TABLE_SUMMARY, 100)))
    summary_table = next(
        table for table in document.tables
        if [cell.text for cell in table.rows[0].cells] == ["项目", "说明"]
    )
    table_text = [[cell.text for cell in row.cells] for row in summary_table.rows]
    assert table_text == [["项目", "说明"], ["问题", "SQL injection"], ["来源", "CNVD"]]

    elements = _build_summary_markdown_elements(TABLE_SUMMARY, _build_styles(_ensure_chinese_font()))
    table = next(element for element in elements if element.__class__.__name__ == "Table")
    pdf_table_text = [
        cell.getPlainText()
        for row in table._cellvalues
        for cell in row
        if hasattr(cell, "getPlainText")
    ]
    assert pdf_table_text == ["项目", "说明", "问题", "SQL injection", "来源", "CNVD"]


def test_domain_report_exports_keep_markdown_tables_as_native_tables():
    html, html_mime = export_markdown_report("领域报告", DOMAIN_TABLE_REPORT, "html")
    word, word_mime = export_markdown_report("领域报告", DOMAIN_TABLE_REPORT, "word")
    pdf, pdf_mime = export_markdown_report("领域报告", DOMAIN_TABLE_REPORT, "pdf")

    assert html_mime == "text/html"
    html_text = html.decode("utf-8") if isinstance(html, bytes) else html
    assert '<table class="report-summary-table">' in html_text
    assert ".report-summary-table{width:100%;table-layout:fixed" in html_text
    assert "overflow-wrap:anywhere" in html_text
    assert "<th>项目</th>" in html_text and "<td>未执行</td>" in html_text
    assert "| --- | --- | --- |" not in html_text

    assert word_mime.endswith("wordprocessingml.document")
    document = Document(BytesIO(word))
    native_table = next(table for table in document.tables if table.rows[0].cells[0].text == "项目")
    assert [[cell.text for cell in row.cells] for row in native_table.rows] == [
        ["项目", "观测值", "判定"],
        ["黑盒状态", "未执行", "本次只请求白盒验证"],
        [
            "超长证据",
            "长证据内容用于确认 PDF 和 Word 会自动换行，而不是把表格竖线当作正文输出。",
            "已记录",
        ],
    ]
    assert not any("| --- |" in paragraph.text for paragraph in document.paragraphs)

    assert pdf_mime == "application/pdf"
    assert pdf.startswith(b"%PDF-")
    blocks = markdown_to_blocks(DOMAIN_TABLE_REPORT)
    table_block = next(block for block in blocks if block.kind == "table")
    assert len(table_block.rows) == 3
    styles = _build_styles(_ensure_chinese_font())
    elements = []
    _append_markdown_elements(elements, DOMAIN_TABLE_REPORT, styles)
    pdf_table = next(element for element in elements if isinstance(element, Table))
    assert pdf_table.repeatRows == 1
    assert pdf_table._colWidths == [(A4[0] - 36 * mm) / 3] * 3
    assert [cell.getPlainText() for row in pdf_table._cellvalues for cell in row] == [
        "项目", "观测值", "判定", "黑盒状态", "未执行", "本次只请求白盒验证", "超长证据",
        "长证据内容用于确认 PDF 和 Word 会自动换行，而不是把表格竖线当作正文输出。",
        "已记录",
    ]


def test_domain_word_export_uses_compact_title_for_long_task_names():
    title = "沙箱黑白盒测试 · sbx_3b317e05c0fa45128c20e31d"
    word, content_type = export_markdown_report(title, DOMAIN_TABLE_REPORT, "word")

    assert content_type.endswith("wordprocessingml.document")
    document = Document(BytesIO(word))
    title_paragraph = document.paragraphs[0]
    assert title_paragraph.text == "沙箱黑白盒测试"
    assert title_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title_paragraph.runs
    assert all(run.font.size is not None and run.font.size.pt == 18 for run in title_paragraph.runs)
    assert document.paragraphs[1].text == "任务编号：sbx_3b317e05c0fa45128c20e31d"
    assert document.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_pdf_title_puts_long_task_id_on_a_separate_wrappable_line():
    title = "沙箱黑白盒测试 · sbx_3b317e05c0fa45128c20e31d"

    markup = domain_report_exporter._pdf_title_markup(title)

    assert "沙箱黑白盒测试" in markup
    assert "任务编号" in markup
    assert "<br/>" in markup
    assert "sbx_3b317e05c0fa45128c20e31d" in markup
    assert "sbx_3b317e05c0fa<br/>" not in markup


def test_pdf_markdown_keeps_project_names_and_sandbox_ids_readable():
    styles = _build_styles(_ensure_chinese_font())
    elements = []
    _append_markdown_elements(
        elements,
        "**3.- 位置：侦察阶段\n\n任务：沙箱测试 · sbx_3b317e05c0fa45128c20e31d\n\n表达式 a*b 和 glob *.py",
        styles,
    )

    rendered_text = "\n".join(element.getPlainText() for element in elements if hasattr(element, "getPlainText"))
    project_name = "Local-File-Inclusion-Demo"
    sandbox_id = "sbx_3b317e05c0fa45128c20e31d"
    markup = reportlab_markup((MarkdownSegment(f"项目：{project_name} 任务编号：{sandbox_id}"),))
    long_digest_markup = reportlab_markup((MarkdownSegment(f"sha256:{'a' * 64}"),))

    assert "**" not in rendered_text
    assert "a*b" in rendered_text
    assert "*.py" in rendered_text
    assert project_name in markup
    assert sandbox_id in markup
    assert "Local-File-Inclusion-Dem<br/>o" not in markup
    assert "sbx_3b317e05c0fa<br/>45128c20e31d" not in markup
    assert "<br/>" in long_digest_markup


def test_pdf_final_caveat_stays_with_the_previous_evidence_item():
    styles = _build_styles(_ensure_chinese_font())
    elements = []

    _append_markdown_elements(
        elements,
        "- 最后一项证据仍待复核。\n\n> 说明：本次未采集响应头，不能据此确认相关问题。",
        styles,
    )

    assert isinstance(elements[-1], KeepTogether)


def test_domain_coverage_names_requested_mode_and_does_not_claim_all_ranges_were_scanned():
    lines = _coverage_lines({
        "verification_status": "complete",
        "coverage_status": "complete",
        "requested_mode": "whitebox",
        "deterministic_status": "passed",
        "ai_dynamic_status": "skipped",
        "blackbox_status": "not_required",
        "remote_blackbox_status": "not_required",
        "unscanned": [],
    })

    joined = "\n".join(lines)
    assert "验证模式：白盒" in joined
    assert "确定性检查：已通过" in joined
    assert "AI 动态用例：未执行" in joined
    assert "黑盒验证：未纳入本次验证" in joined
    assert "远程黑盒验证：未纳入本次验证" in joined
    assert "回执未列出具体未扫描项" in joined
    assert "已记录范围内无遗漏" not in joined


def test_domain_coverage_keeps_coverage_and_verification_status_distinct():
    joined = "\n".join(_coverage_lines({
        "coverage_status": "partial",
        "verification_status": "complete",
    }))

    assert "覆盖状态：部分覆盖" in joined
    assert "验证状态：核验已完成" in joined

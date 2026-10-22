"""基于领域原始结果生成沙箱、渗透与远程评估报告。"""

from __future__ import annotations

import json
import re
from html import escape
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.services.report_exporter import _sanitize_report_html
from app.services.report_markdown import markdown_to_html
from app.services.report_pdf_exporter import (
    _append_markdown_elements,
    _build_styles,
    _ensure_chinese_font,
)
from app.services.report_text import clean_report_markdown, clean_report_payload, clean_report_text
from app.services.report_word_exporter import _append_markdown_to_word, _set_run_chinese_font


def _value(value: Any, fallback: str = "未记录") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list)):
        return json.dumps(clean_report_payload(value), ensure_ascii=False, indent=2, default=str)
    return clean_report_text(value)


def _status_label(value: Any) -> str:
    labels = {
        "success": "执行完成",
        "succeeded": "执行完成",
        "completed": "执行完成",
        "complete": "执行完成",
        "verified": "已核验通过",
        "passed": "已通过",
        "partial": "部分完成",
        "failed": "执行失败",
        "skipped": "未执行",
        "unknown": "未知",
        "not_required": "不适用",
        "receipt_missing": "缺少有效回执",
        "cancelled": "已取消",
        "stopped": "已停止",
        "expired": "已过期",
        "running": "执行中",
        "pending": "等待执行",
    }
    return labels.get(str(value or ""), str(value or "状态未知"))


def _verification_status_label(value: Any) -> str:
    labels = {
        "complete": "核验已完成",
        "completed": "核验已完成",
        "verified": "核验通过",
        "passed": "核验通过",
        "partial": "部分核验",
        "failed": "未通过",
        "unknown": "未知",
        "pending": "等待核验",
        "running": "核验中",
        "cancelled": "已取消",
        "stopped": "已停止",
        "expired": "已过期",
        "receipt_missing": "缺少有效回执",
    }
    return labels.get(str(value or ""), _status_label(value))


def _localize_sandbox_report_statuses(markdown: str) -> str:
    """只翻译报告正文中的状态字段，不改写代码块或技术证据。"""
    fields = (
        "确定性执行结论", "确定性执行", "确定性检查", "AI 动态用例", "黑盒路由",
        "系统验证门禁", "最终状态", "验证状态",
    )
    field_pattern = "|".join(re.escape(field) for field in fields)
    status_pattern = (
        r"succeeded|success|passed|completed|complete|partial|failed|skipped|unknown|"
        r"not_required|receipt_missing|running|pending|stopped|expired|cancelled"
    )
    expression = re.compile(
        rf"(?P<prefix>(?:{field_pattern})\s*[=:：]\s*)"
        rf"(?P<status>{status_pattern})(?=$|[\s，,。；;])",
        re.I,
    )
    lines: list[str] = []
    fence_char: str | None = None
    fence_length = 0
    for line in markdown.splitlines(keepends=True):
        marker = re.match(r"^ {0,3}(`{3,}|~{3,})(.*)$", line.rstrip("\r\n"))
        if marker:
            token, remainder = marker.group(1), marker.group(2)
            if fence_char is None:
                if token[0] != "`" or "`" not in remainder:
                    fence_char, fence_length = token[0], len(token)
            elif token[0] == fence_char and len(token) >= fence_length and not remainder.strip():
                fence_char, fence_length = None, 0
            lines.append(line)
            continue
        if fence_char is None:
            line = expression.sub(lambda match: f"{match.group('prefix')}{_status_label(match.group('status'))}", line)
        lines.append(line)
    return "".join(lines)


def _coverage_lines(coverage: Any) -> list[str]:
    if not isinstance(coverage, dict):
        return ["覆盖范围：未记录"]
    coverage_status = coverage.get("coverage_status")
    verification_status = coverage.get("verification_status")
    coverage_labels = {
        "complete": "完整覆盖",
        "completed": "完整覆盖",
        "verified": "完整覆盖",
        "partial": "部分覆盖",
        "incomplete": "覆盖不完整",
        "failed": "覆盖失败",
        "unknown": "未知",
        "cancelled": "已取消",
    }
    lines = []
    if coverage_status is not None:
        lines.append(f"覆盖状态：{coverage_labels.get(str(coverage_status), _status_label(coverage_status))}")
    elif verification_status is not None:
        # 旧报告只有验证状态时，以其作为覆盖状态的唯一可用事实。
        lines.append(f"覆盖状态：{coverage_labels.get(str(verification_status), _status_label(verification_status))}")
    else:
        lines.append("覆盖状态：未知")
    if verification_status is not None:
        lines.append(f"验证状态：{_verification_status_label(verification_status)}")
    requested_mode = str(coverage.get("requested_mode") or "").strip().lower()
    if requested_mode:
        mode_labels = {"whitebox": "白盒", "blackbox": "黑盒", "combined": "黑白盒", "deploy": "部署检查"}
        lines.append(f"验证模式：{mode_labels.get(requested_mode, requested_mode)}")
    for key, label in (
        ("deterministic_status", "确定性检查"),
        ("ai_dynamic_status", "AI 动态用例"),
        ("blackbox_status", "黑盒验证"),
        ("remote_blackbox_status", "远程黑盒验证"),
    ):
        if coverage.get(key) is not None:
            status_label = "未纳入本次验证" if coverage[key] == "not_required" else _status_label(coverage[key])
            lines.append(f"{label}：{status_label}")
    for key, label in (
        ("scan_mode", "扫描模式"),
        ("total_file_count", "目标文件数"),
        ("scanned_file_count", "已扫描文件数"),
        ("skipped_file_count", "跳过文件数"),
        ("coverage_ratio", "文件覆盖比例"),
        ("semantic_char_coverage_ratio", "语义字符覆盖比例"),
    ):
        if coverage.get(key) is not None:
            lines.append(f"{label}：{_value(coverage[key])}")
    unscanned = coverage.get("unscanned")
    if isinstance(unscanned, list) and unscanned:
        lines.append("未扫描或未验证范围：")
        lines.extend(f"{index}. {_value(item)}" for index, item in enumerate(unscanned, 1))
    else:
        lines.append("未扫描或未验证范围：回执未列出具体未扫描项")
    if requested_mode:
        lines.append("覆盖状态只适用于上述请求模式及已记录阶段；其他模式不自动视为已执行")
    skipped = coverage.get("skipped")
    if isinstance(skipped, list) and skipped:
        lines.append("跳过项目：")
        lines.extend(f"{index}. {_value(item)}" for index, item in enumerate(skipped, 1))
    return lines


def score_section_lines(
    score: Any,
    *,
    available: bool,
    scope_label: str = "",
    reason: str = "",
    rule: str = "",
    audit_label: str = "",
    finding_count: int | None = None,
    severity_counts: dict[str, Any] | None = None,
    score_breakdown: dict[str, Any] | None = None,
) -> list[str]:
    """导出正文统一的“评分与结论”：出分写明范围与依据，不出分写明原因，绝不用占位分。"""
    lines = ["## 评分与结论"]
    is_cvss_policy = isinstance(score_breakdown, dict) and score_breakdown.get("policy_version") == "cvss-weighted-v1"
    if available and isinstance(score, (int, float)) and not isinstance(score, bool):
        label = (
            str(score_breakdown.get("score_label") or "平台综合分（非 CVSS 官方综合分）")
            if is_cvss_policy
            else "评分"
        )
        lines.append(f"{label}：{int(score)} / 100")
        if scope_label:
            lines.append(f"评分范围：{scope_label}")
        if rule:
            lines.append(f"计分规则：{rule}")
        if audit_label:
            lines.append(f"复核状态：{audit_label}")
        if is_cvss_policy:
            components = score_breakdown.get("components") or {}
            cvss = components.get("cvss") if isinstance(components, dict) else {}
            cvss = cvss if isinstance(cvss, dict) else {}
            vectors = score_breakdown.get("cvss_vectors")
            vectors = vectors if isinstance(vectors, list) else []
            standard = score_breakdown.get("cvss_standard") or "CVSS 3.1（FIRST）"
            lines.append(f"CVSS 来源：{_value(standard)}")
            lines.append(
                "CVSS 向量：" + ("、".join(str(item) for item in vectors) if vectors else "无有效向量")
            )
            cvss_weight = cvss.get("weight", 0)
            lines.append(f"CVSS 权重：{int(round(cvss_weight * 100))}%" if cvss_weight else "CVSS 权重：0%")
            max_base_score = cvss.get("max_base_score")
            if max_base_score is not None:
                lines.append(f"CVSS 最高基础分：{_value(max_base_score)} / 10")
            else:
                lines.append("CVSS 最高基础分：未评估")
            if score_breakdown.get("cvss_reason"):
                lines.append(f"CVSS 说明：{_value(score_breakdown['cvss_reason'])}")
    else:
        lines.append("评分：未评估")
        lines.append(f"不出分原因：{reason or '缺少可计分的已测范围'}")
    if finding_count is not None:
        counts = {
            str(key): int(value) for key, value in (severity_counts or {}).items()
            if isinstance(value, int) and not isinstance(value, bool) and value > 0
        }
        detail = "、".join(f"{key} {value} 项" for key, value in counts.items())
        if finding_count:
            lines.append(
                f"结论：已测范围内记录 {finding_count} 项发现"
                + (f"（{detail}）" if detail else "") + "，请按修复建议逐项处理并复测。"
            )
        else:
            lines.append("结论：已测范围内未记录发现；这只代表本次已测范围，不等于已证明不存在风险。")
    return lines


def insert_section_after_task_info(markdown: str, section: list[str]) -> str:
    """把评分段放在“任务信息”之后、正文之前，阅读时先看到结论。"""
    text = markdown.rstrip("\n")
    marker = "\n## "
    first = text.find(marker)
    second = text.find(marker, first + 1) if first >= 0 else -1
    block = "\n".join(section)
    if second < 0:
        return f"{text}\n\n{block}\n"
    return f"{text[:second]}\n\n{block}\n{text[second:]}\n"


def build_pentest_report_markdown(
    domain_data: dict[str, Any],
    *,
    task_name: str,
    project_name: str,
    execution_status: str,
    error_message: str | None,
    coverage: dict[str, Any],
) -> str:
    """把授权渗透任务的真实阶段、发现、证据和建议写入可读报告。"""
    scope = domain_data.get("scope") if isinstance(domain_data.get("scope"), dict) else {}
    authorized_target = (
        scope.get("remote_target_url")
        or scope.get("sandbox_env_public_id")
        or domain_data.get("target_type_label")
    )
    lines = [
        "# 授权渗透测试报告",
        "",
        "## 任务信息",
        f"项目：{project_name or '未关联项目'}",
        f"任务：{task_name or domain_data.get('task_name') or '未命名任务'}",
        f"执行状态：{_status_label(execution_status)}",
        f"报告状态：{_status_label(domain_data.get('status'))}",
        f"授权目标：{_value(authorized_target)}",
        f"授权窗口：{_value(domain_data.get('window_start'))} 至 {_value(domain_data.get('window_end'))}",
        "",
        "## 覆盖范围",
        *_coverage_lines(coverage),
    ]
    if scope.get("notes"):
        lines.append(f"范围补充：{_value(scope['notes'])}")
    if error_message:
        lines.extend(["", "## 执行异常", _value(error_message)])

    phases = domain_data.get("phases") if isinstance(domain_data.get("phases"), list) else []
    lines.extend(["", "## 执行阶段"])
    if not phases:
        lines.append("没有保存可核对的阶段记录。")
    for phase in phases:
        if not isinstance(phase, dict):
            continue
        lines.extend(
            [
                "",
                f"### {_value(phase.get('phase_label') or phase.get('phase_key'), '未命名阶段')}",
                f"状态：{_status_label(phase.get('status'))}",
            ]
        )
        if phase.get("summary"):
            lines.append(f"阶段结论：{_value(phase['summary'])}")
        if phase.get("error_message"):
            lines.append(f"阶段错误：{_value(phase['error_message'])}")
        for line in phase.get("lines", []) if isinstance(phase.get("lines"), list) else []:
            if isinstance(line, dict):
                lines.append(f"推演线：{_value(line.get('hypothesis'))}；状态：{_status_label(line.get('status'))}")

    findings = domain_data.get("findings") if isinstance(domain_data.get("findings"), list) else []
    lines.extend(["", "## 发现清单", f"共记录 {len(findings)} 项。"])
    if not findings:
        lines.append("本次范围未保存可报告发现；这不等于已证明不存在风险。")
    for index, finding in enumerate(findings, 1):
        if not isinstance(finding, dict):
            continue
        lines.extend(
            [
                "",
                f"### {index}. {_value(finding.get('title'), '未命名发现')}",
                f"严重度：{_value(finding.get('severity'))}",
                f"复核状态：{_value(finding.get('status'))}",
                f"问题说明：{_value(finding.get('description'))}",
            ]
        )
        if finding.get("repro_steps"):
            lines.extend(["复现步骤：", _value(finding["repro_steps"])])
        if finding.get("evidence"):
            lines.extend(["证据：", "```json", _value(finding["evidence"]), "```"])
        if finding.get("fix_advice"):
            lines.extend(["修复建议：", _value(finding["fix_advice"])])
    return "\n".join(lines).strip() + "\n"


def build_sandbox_report_markdown(
    domain_data: dict[str, Any],
    *,
    task_name: str,
    project_name: str,
) -> str:
    """保留沙箱原始报告正文，并补齐任务、状态和覆盖元数据。"""
    original = _localize_sandbox_report_statuses(
        clean_report_markdown(str(domain_data.get("report_md") or "").strip())
    )
    if not original:
        original = (
            "原始测试报告正文缺失，无法还原逐项测试内容与结果。"
            "以下仅展示任务状态和结构化覆盖信息，不代表已验证具体测试项。"
        )
    coverage = domain_data.get("coverage") if isinstance(domain_data.get("coverage"), dict) else {}
    lines = [
        "# 沙箱测试报告",
        "",
        "## 任务信息",
        f"项目：{project_name or '未关联项目'}",
        f"任务：{task_name or '未命名任务'}",
        f"执行状态：{_status_label(domain_data.get('execution_status'))}",
        "验证状态：" + _verification_status_label(
            domain_data.get("verification_status") or coverage.get("verification_status")
        ),
        "",
        "## 覆盖范围",
        *_coverage_lines(coverage),
    ]
    if domain_data.get("error_message"):
        lines.extend(["", "## 执行异常", _value(domain_data["error_message"])])
    lines.extend(["", "## 原始测试报告", original])
    remote = (
        domain_data.get("evidence", {}).get("remote_blackbox")
        if isinstance(domain_data.get("evidence"), dict)
        else None
    )
    if isinstance(remote, dict):
        lines.extend(
            [
                "",
                "## 远程目标有限黑盒检查",
                "本次最多进行有限次数的 HTTPS 根路径只读请求；未保存目标响应正文，也未进行完整主动渗透。",
            ]
        )
        lines.append(f"授权目标：{_value(remote.get('target_origin'))}")
        lines.append("请求方法：GET")
        for key, label in (
            ("resolved_ip", "解析地址"),
            ("status_code", "HTTP 状态码"),
            ("latency_ms", "响应耗时毫秒"),
            ("redirect_location", "重定向目标"),
            ("content_type", "内容类型"),
            ("security_headers", "安全响应头"),
            ("missing_security_headers", "缺失安全响应头"),
            ("body_bytes_sampled", "读取响应字节数"),
        ):
            if remote.get(key) is not None:
                lines.append(f"{label}：{_value(remote[key])}")
    return "\n".join(lines).strip() + "\n"


def build_remote_pentest_report_markdown(content: dict[str, Any]) -> str:
    """把远程网站评估的授权快照、七阶段证据、限制和发现完整写入报告。"""
    scope = content.get("scope") if isinstance(content.get("scope"), dict) else {}
    methodology = content.get("methodology") if isinstance(content.get("methodology"), dict) else {}
    lines = [
        "# 远程网站安全评估报告",
        "",
        "## 任务信息",
        f"任务：{_value(scope.get('task_name'))}",
        f"目标：{_value(scope.get('target_url'))}",
        f"状态：{_status_label(content.get('status'))}",
        f"规则版本：{_value(scope.get('rules_version'))}",
        f"授权窗口：{_value(scope.get('window_start'))} 至 {_value(scope.get('window_end'))}",
        f"授权依据：{_value(scope.get('authorization_basis'))}",
        f"范围补充：{_value(scope.get('scope_notes'))}",
        "",
        "## 方法与覆盖边界",
        f"请求方式：{_value(methodology.get('request_method'))}",
        f"请求上限：{_value(methodology.get('request_count_max'))}",
        f"目标范围：{_value(methodology.get('target_scope'))}",
        f"响应正文保存：{'是' if methodology.get('response_body_saved') else '否'}",
        f"覆盖状态：{_value(methodology.get('coverage_note') or methodology.get('coverage_state'))}",
    ]
    limitations = methodology.get("limitations") if isinstance(methodology.get("limitations"), list) else []
    if limitations:
        lines.append("限制：")
        lines.extend(f"{index}. {_value(item)}" for index, item in enumerate(limitations, 1))
    lines.extend(["", "## 阶段记录"])
    phases = content.get("phases") if isinstance(content.get("phases"), list) else []
    for phase in phases:
        if not isinstance(phase, dict):
            continue
        lines.extend(
            ["", f"### {_value(phase.get('label') or phase.get('key'))}", f"状态：{_status_label(phase.get('status'))}"]
        )
        if phase.get("summary"):
            lines.append(f"阶段说明：{_value(phase['summary'])}")
        if phase.get("evidence"):
            lines.extend(["阶段证据：", "```json", _value(phase["evidence"]), "```"])
        if phase.get("error_message"):
            lines.append(f"阶段错误：{_value(phase['error_message'])}")
    findings = content.get("findings") if isinstance(content.get("findings"), list) else []
    lines.extend(["", "## 配置观察", f"共记录 {len(findings)} 项。"])
    if not findings:
        lines.append("本次有限检查未记录配置观察；这不等于目标已通过完整安全测试。")
    for index, finding in enumerate(findings, 1):
        if not isinstance(finding, dict):
            continue
        lines.extend(
            [
                "",
                f"### {index}. {_value(finding.get('title'), '未命名观察')}",
                f"严重度：{_value(finding.get('severity'))}",
                f"状态：{_value(finding.get('status'))}",
                f"说明：{_value(finding.get('description'))}",
            ]
        )
        if finding.get("evidence"):
            lines.extend(["证据：", "```json", _value(finding["evidence"]), "```"])
        if finding.get("fix_advice"):
            lines.append(f"修复建议：{_value(finding['fix_advice'])}")
    failure = content.get("failure")
    if isinstance(failure, dict):
        lines.extend(
            ["", "## 失败信息", f"错误代码：{_value(failure.get('code'))}", f"说明：{_value(failure.get('message'))}"]
        )
    markdown = "\n".join(lines).strip() + "\n"
    score = content.get("score") if isinstance(content.get("score"), dict) else {}
    markdown = insert_section_after_task_info(
        markdown,
        score_section_lines(
            score.get("score"),
            available=bool(score.get("available")),
            scope_label=str(score.get("scope_label") or ""),
            reason=str(score.get("reason") or "报告未保存评分"),
            rule=str(score.get("rule") or ""),
            audit_label=str(score.get("audit_label") or ""),
            finding_count=sum(
                1 for item in findings
                if isinstance(item, dict) and item.get("status") == "config_confirmed"
            ),
            severity_counts=(
                score.get("severity_counts")
                if isinstance(score.get("severity_counts"), dict)
                else None
            ),
            score_breakdown=score,
        ),
    )
    return clean_report_markdown(markdown)


def export_markdown_report(title: str, markdown: str, format: str) -> tuple[bytes, str]:
    """将同一份报告正文转换为可读 HTML、PDF 或 Word。"""
    if format == "html":
        body = markdown_to_html(markdown)
        html = (
            '<style>.report-summary-table{width:100%;table-layout:fixed;border-collapse:collapse;}'
            '.report-summary-table th,.report-summary-table td{border:1px solid #d1d5db;padding:8px;'
            'text-align:left;vertical-align:top;overflow-wrap:anywhere;word-break:break-word;}'
            '.report-summary-table th{background:#e8eef8;}'
            'pre{max-width:100%;box-sizing:border-box;overflow-x:auto;white-space:pre-wrap;'
            'overflow-wrap:anywhere;word-break:break-word;}'
            'pre code{white-space:inherit;overflow-wrap:inherit;word-break:inherit;}</style>'
            '<main style="box-sizing:border-box;width:100%;max-width:960px;margin:32px auto;padding:32px;'
            'overflow-wrap:anywhere;'
            'font:16px/1.7 -apple-system,BlinkMacSystemFont,Segoe UI,sans-serif;color:#172033">'
            f"<h1>{escape(title)}</h1>{body}</main>"
        )
        return _sanitize_report_html(html).encode("utf-8"), "text/html"
    if format == "pdf":
        font = _ensure_chinese_font()
        styles = _build_styles(font)
        buffer = BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            title=title,
        )
        elements = [Paragraph(_pdf_title_markup(title), styles["title"]), Spacer(1, 5 * mm)]
        _append_markdown_elements(elements, markdown, styles)
        document.build(elements)
        return buffer.getvalue(), "application/pdf"
    if format == "word":
        document = Document()
        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_paragraph.paragraph_format.space_after = Pt(3)
        title_paragraph.paragraph_format.keep_with_next = True
        title_label, public_id = _split_report_title(title)
        title_run = title_paragraph.add_run(title_label)
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(0x17, 0x20, 0x33)
        _set_run_chinese_font(title_run, "SimSun")
        if public_id:
            identifier_paragraph = document.add_paragraph()
            identifier_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            identifier_paragraph.paragraph_format.space_after = Pt(10)
            identifier_run = identifier_paragraph.add_run(f"任务编号：{public_id}")
            identifier_run.font.size = Pt(9)
            identifier_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            _set_run_chinese_font(identifier_run, "SimSun")
        else:
            title_paragraph.paragraph_format.space_after = Pt(16)
        _append_markdown_to_word(document, markdown)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raise ValueError(f"不支持的领域报告格式：{format}")


def _pdf_title_markup(title: str) -> str:
    """将带任务编号的长标题拆成可读、可换行的 PDF 标题。"""
    label, public_id = _split_report_title(title)
    if public_id:
        return (
            f"{escape(label)}<br/>"
            f'<font size="9" color="#64748B">任务编号：{escape(public_id)}</font>'
        )
    return escape(title)


def _split_report_title(title: str) -> tuple[str, str | None]:
    """Keep a long machine identifier out of the human-readable report title."""
    label, separator, public_id = title.rpartition(" · ")
    if separator and len(public_id) >= 20 and re.fullmatch(r"[A-Za-z0-9_-]+", public_id):
        return label, public_id
    return title, None
